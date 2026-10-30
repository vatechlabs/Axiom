"""File readers and per-client reconciliation. Never confuse source date with import date."""
import hashlib, io, json, os, re
from datetime import date,datetime
from pathlib import Path
from django.conf import settings
from django.db import transaction
from .models import (Claim,ImportBatch,LagClaimSnapshot,ErrorObservation,ErrorHistory,
                     EdiFile,EdiClaimIndex,EdiDocument)
from openpyxl import load_workbook

SHEETS=('LOCAL CLMS','LOCAL ADJ','ITS CLMS','ITS ADJ')
# Explicit per-column overrides for real workbook (not guessed from encrypted source).
# Example: {"DCN":"B", "CLAIM_SUBMITTER_ID":"C", "TPA_AGE":"Q", "TOTAL_AGE":"R"}
COLS=json.loads(os.environ.get('LAG_COLUMN_MAP_JSON','{}'))
ALIASES={'dcn':['DCN','DOCUMENT CONTROL NUMBER'],
         'submitter':['CLAIM SUBMITTER ID','CLAIM NUMBER','CLAIM #','CLAIM NO','SUBMITTER ID'],
         'tpa_age':['TPA AGE','TPA DAYS','AGE WITH TPA'],
         'total_age':['TOTAL AGE','OVERALL AGE','AGE','TOTAL DAYS'],
         'tpa':['TPA','PAYER','ADMINISTRATOR']}
def canonical(s):return re.sub(r'[^A-Z0-9]+',' ',str(s or '').upper()).strip()
def column_index(name,headers):
    override=COLS.get(name.upper())
    if override:
        from openpyxl.utils.cell import column_index_from_string
        return column_index_from_string(override.upper())-1
    for alias in ALIASES.get(name,[]):
        for j,v in enumerate(headers):
            if canonical(v)==canonical(alias):return j
    return None

def is_encrypted(blob):return blob.startswith(bytes.fromhex('D0CF11E0A1B11AE1'))
def workbook_from_bytes(blob,password=None):
    if is_encrypted(blob):
        if not password:raise ValueError('This workbook is encrypted. Enter its password in the upload form.')
        try:import msoffcrypto
        except ImportError as exc:raise ValueError('msoffcrypto-tool is not installed.') from exc
        out=io.BytesIO()
        f=msoffcrypto.OfficeFile(io.BytesIO(blob));f.load_key(password=password)
        try:f.decrypt(out)
        except Exception as exc:raise ValueError('Unable to decrypt workbook: check password.') from exc
        blob=out.getvalue()
    return load_workbook(io.BytesIO(blob),read_only=True,data_only=True)

def validate_lag_client(client, blob, password=None):
    """Reject a known client mismatch before duplicate checks or any database writes.

    Deliberately do not infer a client from filename, fund, or claim identifiers.
    Unknown workbook layouts are rejected for Highmark until its distinct SCCF
    importer is implemented; otherwise a WGS workbook could corrupt its claims.
    """
    wb = workbook_from_bytes(blob, password)
    try:
        sheets = {canonical(s) for s in wb.sheetnames}
        wgs_sheets = {canonical(s) for s in SHEETS}
        # WGS-specific four-tab LAG layout, which the current parser expects.
        is_wgs = wgs_sheets.issubset(sheets)
        is_highmark = False
        for ws in wb.worksheets[:5]:
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
                headers = {canonical(v) for v in row if v is not None}
                if 'SCCF' in headers and ('CAID' in headers or 'SCCF DATE' in headers):
                    is_highmark = True
                    break
            if is_highmark:
                break
    finally:
        wb.close()
    code = client.code.upper()
    if code in ('HIGHMARK', 'HMRK', 'HM'):
        if is_wgs:
            raise ValueError('Wrong LAG report: this is a WGS LAG report. Please select WGS, not Highmark.')
        if not is_highmark:
            raise ValueError('Highmark requires a Highmark LAG report with SCCF/CAID columns. This file could not be verified as Highmark.')
    elif code == 'WGS' and is_highmark and not is_wgs:
        raise ValueError('Wrong LAG report: this is a Highmark LAG report. Please select Highmark, not WGS.')

def parse_lag(blob,password=None):
    wb=workbook_from_bytes(blob,password)
    found=[s for s in SHEETS if s in wb.sheetnames]
    if len(found)!=4:raise ValueError('Expected four sheets: '+', '.join(SHEETS)+'. Found: '+', '.join(wb.sheetnames))
    rows=[]
    for sheet in found:
        ws=wb[sheet]
        vals=ws.iter_rows(values_only=True)
        headerrow=None; mapped=None
        for pos,row in enumerate(vals,1):
            if pos>30:break
            headers=list(row)
            fields={key:column_index(key,headers) for key in ('dcn','submitter','tpa_age','total_age','tpa')}
            if (fields['dcn'] is not None or fields['submitter'] is not None) and fields['tpa_age'] is not None and fields['total_age'] is not None:
                headerrow=pos;mapped=fields;break
        if headerrow is None:raise ValueError(f'{sheet}: claim and age headers not detected. Configure LAG_COLUMN_MAP_JSON using actual columns.')
        for pos,row in enumerate(vals,headerrow+1):
            def get(key):
                idx=mapped[key];return row[idx] if idx is not None and idx<len(row) else None
            def ident(v):
                if v is None:return ''
                if isinstance(v,float) and v.is_integer():return str(int(v))
                return str(v).strip()
            dcn=ident(get('dcn'));submitter=ident(get('submitter'))
            if not (dcn or submitter):continue
            def age(x):
                if x is None or str(x).strip()=='':return 0
                v=int(float(x));
                if v<0:raise ValueError('Negative aging value')
                return v
            try:tpa_age=age(get('tpa_age'));total_age=age(get('total_age'))
            except Exception as exc:raise ValueError(f'{sheet} row {pos}: invalid age ({exc})') from exc
            # Confirmed with user: Fund = Excel column I; do NOT use client code here.
            fund=ident(row[8]) if len(row)>8 else ''
            raw={f'col_{i+1}':v.isoformat() if isinstance(v,(date,datetime)) else str(v or '') for i,v in enumerate(row)}
            rows.append({'sheet':sheet,'row':pos,'dcn':dcn,'claim_number':submitter or dcn,
                         'fund':fund,'tpa_age':tpa_age,'total_age':total_age,'tpa':ident(get('tpa')),'raw':raw})
    return rows

def parse_highmark_lag(blob, password=None):
    """Read a verified Highmark SCCF/CAID workbook without using WGS sheet rules.

    SCCF identifies the claim; CAID is retained as the secondary identifier.
    Reject missing IDs or conflicting repeated SCCFs. Do not invent aging values:
    blank/unavailable aging maps to 0 (the current claim model requires integers).
    """
    wb = workbook_from_bytes(blob, password)
    rows = []
    seen = {}
    try:
        for ws in wb.worksheets:
            header_position = None
            columns = None
            for number, values in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True), 1):
                hdr = {canonical(value): index for index, value in enumerate(values) if value is not None}
                if 'SCCF' in hdr and ('CAID' in hdr or 'SCCF DATE' in hdr):
                    header_position = number
                    columns = hdr
                    break
            if header_position is None:
                continue

            def pick(values, *names):
                for name in names:
                    at = columns.get(name)
                    if at is not None and at < len(values) and values[at] is not None:
                        return values[at]
                return None

            def identifier(value):
                if value is None:
                    return ''
                if isinstance(value, float) and value.is_integer():
                    return str(int(value))
                return str(value).strip()

            def aging(value, location):
                if value is None or str(value).strip() == '':
                    return 0
                try:
                    integer = int(float(value))
                except (TypeError, ValueError) as exc:
                    raise ValueError(f'{location}: invalid aging value') from exc
                if integer < 0:
                    raise ValueError(f'{location}: negative aging value')
                return integer

            for row_number, values in enumerate(ws.iter_rows(min_row=header_position + 1, values_only=True), header_position + 1):
                claim_id = identifier(pick(values, 'SCCF'))
                caid = identifier(pick(values, 'CAID'))
                if not claim_id and not caid:
                    continue
                if not claim_id:
                    raise ValueError(f'{ws.title} row {row_number}: SCCF missing; refusing to guess the claim identifier')
                if len(claim_id) > 80 or len(caid) > 80:
                    raise ValueError(f'{ws.title} row {row_number}: claim identifier exceeds database limit')
                if claim_id in seen and seen[claim_id] != caid:
                    raise ValueError(f'{ws.title} row {row_number}: SCCF {claim_id} occurs with conflicting CAID values')
                seen[claim_id] = caid
                tpa_age = aging(pick(values, 'TPA AGE', 'TPA DAYS', 'AGE WITH TPA'), f'{ws.title} row {row_number}')
                total_age = aging(pick(values, 'TOTAL AGE', 'OVERALL AGE', 'TOTAL DAYS'), f'{ws.title} row {row_number}')
                raw = {f'col_{i+1}': v.isoformat() if isinstance(v, (date, datetime)) else str(v if v is not None else '') for i, v in enumerate(values)}
                raw['_highmark_headers'] = {k: str(v) for k, v in columns.items()}
                rows.append({'sheet': ws.title, 'row': row_number, 'dcn': caid,
                             'claim_number': claim_id, 'fund': '', 'tpa': '',
                             'tpa_age': tpa_age, 'total_age': total_age, 'raw': raw})
        if not rows:
            raise ValueError('Highmark workbook has no SCCF claim rows; nothing imported')
        return rows
    finally:
        wb.close()

def import_lag(client,blob,filename,report_date,password=None,user=None):
    from .processing_locks import client_processing_lock, lag_upload_lock, ClientProcessingBusy
    from .models import FolderScanSchedule
    from django.db.models import F
    # Uploads serialize with other LAG writes and all reconciliation, but
    # NOT with the long-running EDI file scan.
    with lag_upload_lock(client.pk):
        schedule,_=FolderScanSchedule.objects.get_or_create(client=client)
        result=_import_lag_unlocked(client,blob,filename,report_date,password,user,
                                    reconcile=False)
        # A scan may be indexing while this upload happens. Persist intent
        # to reconcile, including for duplicate uploads.
        FolderScanSchedule.objects.filter(pk=schedule.pk).update(needs_reconcile=True)
        # If the scan is idle, reconcile immediately; otherwise defer until
        # the scanner finishes. Never block this HTTP request on the scan.
        try:
            with client_processing_lock(client.pk):
                reconcile_client(client)
                FolderScanSchedule.objects.filter(pk=schedule.pk).update(
                    needs_reconcile=False, revision=F('revision')+1)
        except ClientProcessingBusy:
            pass
        return result

def _import_lag_unlocked(client,blob,filename,report_date,password=None,user=None,reconcile=True):
    validate_lag_client(client, blob, password)
    digest=hashlib.sha256(blob+str(report_date).encode('ascii')).hexdigest()
    existing=ImportBatch.objects.defer('original_file').filter(client=client,source='lag',fingerprint=digest,status='processed').first()
    if existing:return existing,True
    rows=(parse_highmark_lag(blob,password) if client.code.upper() in ('HIGHMARK','HMRK','HM')
          else parse_lag(blob,password))
    with transaction.atomic():
        batch=ImportBatch.objects.create(client=client,source='lag',filename=filename[:255],report_date=report_date,
                                        original_file=blob,fingerprint=digest,created_by=user,status='processing')
        for record in rows:
            key=record['claim_number']
            claim,_=Claim.objects.update_or_create(client=client,claim_number=key,
                defaults={'dcn':record['dcn'],'fund':record['fund'],'tpa':record['tpa'],
                          'tpa_age':record['tpa_age'],'total_age':record['total_age'],
                          'report_date':report_date,'source_sheet':record['sheet'],'lag_batch':batch})
            LagClaimSnapshot.objects.create(client=client,batch=batch,claim=claim,row_number=record['row'],
                sheet_name=record['sheet'],row_data={**record['raw'], '_snapshot': {k:record[k] for k in ('dcn','claim_number','fund','tpa','tpa_age','total_age','sheet')}})
        batch.processed_count=len(rows);batch.status='processed';batch.save(update_fields=['processed_count','status'])
    if reconcile:
        reconcile_client(client)
    return batch,False

# Columns from observed fixed-width error report:
# header line contains these labels; boundaries come from dashed line directly beneath.
def parse_error_txt(blob):
    lines=blob.decode('utf-8-sig',errors='replace').splitlines()
    report_date=None
    for line in lines[:12]:
        match=re.search(r'DATE:\s*(\d{2}/\d{2}/\d{2,4})',line)
        if match:report_date=datetime.strptime(match.group(1),'%m/%d/%y' if len(match.group(1).split('/')[-1])==2 else '%m/%d/%Y').date();break
    if not report_date:raise ValueError('Error report DATE header missing')
    # The printed column format is fixed-width, and CLAIM SUBMITTER ID is long.
    # Locate boundaries from the underline ruler rather than split() whitespace.
    results=[];ruler=None
    for i,line in enumerate(lines):
        if line.lstrip().startswith('DCN') and 'ERROR MESSAGE' in line and 'CLAIM SUBMITTER ID' in line:
            if i+1<len(lines):ruler=lines[i+1];break
    if ruler is None:raise ValueError('Unsupported error report layout (missing column header)')
    spans=[(m.start(),m.end()) for m in re.finditer(r'-{2,}',ruler)]
    if len(spans)<8:raise ValueError('Error report ruler could not be parsed')
    # Fixed-width fields by ordinal, aligned to dashed groups in the supplied sample.
    for n,line in enumerate(lines[i+2:],i+3):
        if line.lstrip().startswith('DCN') or line.lstrip().startswith('REPORT NO:') or line.lstrip().startswith('PAGE:'):continue
        dcn=line[spans[0][0]:spans[0][1]].strip()
        if not re.fullmatch(r'[A-Za-z0-9]{6,30}',dcn):continue
        def field(idx):return line[spans[idx][0]:spans[idx][1]].strip() if idx<len(spans) else ''
        results.append({'line':n,'dcn':dcn,'message':field(3),'reject_code':field(4),
                        'claim_submitter_id':field(7),'report_date':report_date})
    return results,report_date

def import_errors(client,path,reconcile=True):
    path=Path(path);blob=path.read_bytes();digest=hashlib.sha256(blob).hexdigest()
    old=ImportBatch.objects.filter(client=client,source='error',fingerprint=digest,status='processed').first()
    if old:
        # Same report moved into Archive: preserve the batch/observations but
        # repair the stored source path once the old location disappears.
        previous = Path(old.file_path) if old.file_path else None
        if previous is None or not previous.is_file():
            ImportBatch.objects.filter(pk=old.pk).update(file_path=str(path), filename=path.name[:255])
            old.file_path = str(path)
            old.filename = path.name[:255]
        return old,True
    rows,report_date=parse_error_txt(blob)
    with transaction.atomic():
        batch=ImportBatch.objects.create(client=client,source='error',filename=path.name[:255],file_path=str(path),
            fingerprint=digest,report_date=report_date,status='processing')
        ErrorObservation.objects.bulk_create([ErrorObservation(client=client,batch=batch,source_line=r['line'],
            dcn=r['dcn'],claim_submitter_id=r['claim_submitter_id'],report_date=report_date,
            message=r['message'],reject_code=r['reject_code']) for r in rows],batch_size=1000)
        batch.processed_count=len(rows);batch.status='processed';batch.save(update_fields=['processed_count','status'])
    if reconcile:
        reconcile_client(client)
    return batch,False

def iter_edi_segments(path, chunk_size=1024*1024):
    """Yield (byte_offset, segment_number, raw_segment) using bounded memory.

    Delimiters come from the fixed-width ISA envelope. Offsets refer to original bytes,
    so a selected claim can later be retrieved using a seek against the remote file.
    """
    with open(path, 'rb') as stream:
        isa = stream.read(106)
        if len(isa) != 106 or not isa.startswith(b'ISA'):
            raise ValueError('Not a valid ISA envelope')
        sep = bytes([isa[105]])
        stream.seek(0)
        pending = b''
        offset = 0
        number = 0
        while True:
            chunk = stream.read(chunk_size)
            if not chunk:
                break
            pending += chunk
            while True:
                end = pending.find(sep)
                if end < 0:
                    break
                segment = pending[:end]
                if segment.strip():
                    yield offset, number, segment
                number += 1
                consumed = end + 1
                offset += consumed
                pending = pending[consumed:]
            if len(pending) > 8*1024*1024:
                raise ValueError('EDI segment exceeds 8 MB')
        if pending.strip():
            raise ValueError('Incomplete trailing EDI segment: file may still be uploading')


def edi_claim_refs(path, kind):
    with open(path, 'rb') as stream:
        isa = stream.read(106)
    if len(isa) != 106 or not isa.startswith(b'ISA'):
        raise ValueError('Not an X12 interchange')
    element = chr(isa[3])
    claim_position = None
    claim_level = False
    seen = set()
    for offset, number, segment in iter_edi_segments(path):
        fields = segment.decode('latin-1').strip().split(element)
        tag = fields[0]
        if kind == '835':
            if tag == 'CLP' and len(fields) > 1:
                if fields[1].strip():
                    yield fields[1].strip(), 'submitter', number, offset
                if len(fields) > 7 and fields[7].strip():
                    yield fields[7].strip(), 'payer_control', number, offset
            continue
        if kind != '837':
            continue
        if tag == 'CLM':
            claim_position = (number, offset)
            claim_level = True
            seen = set()
            if len(fields) > 1 and fields[1].strip():
                yield fields[1].strip(), 'submitter', number, offset
        elif tag in ('ISA', 'GS', 'ST', 'SE', 'GE', 'IEA', 'HL'):
            claim_position = None
            claim_level = False
        elif tag == 'LX':
            claim_level = False
        elif (tag == 'REF' and claim_level and claim_position is not None
              and len(fields) > 2 and fields[1].strip() == '9C'):
            # Verified WGS format; retain the parent CLM position for extraction.
            match = re.fullmatch(r'([0-9]{5}[A-Z]{2}[0-9]{4})-\s*-\s*[0-9]{3}',
                                 fields[2].strip().upper())
            if match and match.group(1) not in seen:
                identifier = match.group(1)
                seen.add(identifier)
                parent_number, parent_offset = claim_position
                yield identifier, 'wgs_ref_9c', parent_number, parent_offset


def parse_edi_claims(blob, kind):
    """Legacy byte-array parser kept for compatibility with existing integrations."""
    import tempfile
    with tempfile.TemporaryFile() as stream:
        stream.write(blob)
        stream.seek(0)
        # Existing callers only need claim ID, identifier type and segment ordinal.
        if len(blob)<106 or not blob.startswith(b'ISA'):
            raise ValueError('Not an X12 interchange')
        element=chr(blob[3]);separator=bytes([blob[105]])
        result=[]
        for number,segment in enumerate(blob.split(separator)):
            fields=segment.decode('latin-1').strip().split(element)
            if len(fields)<2:continue
            if kind=='835' and fields[0]=='CLP':
                if fields[1].strip():result.append((fields[1].strip(),'submitter',number))
                if len(fields)>7 and fields[7].strip():result.append((fields[7].strip(),'payer_control',number))
            if kind=='837' and fields[0]=='CLM' and fields[1].strip():
                result.append((fields[1].strip(),'submitter',number))
        return result


def index_edi(client,path,kind,force=False):
    path=Path(path).resolve()
    before=path.stat()
    old=EdiFile.objects.filter(client=client,kind=kind,path=str(path)).first()
    if not force and old and old.bytes_size==before.st_size and old.modified_ns==before.st_mtime_ns:
        if not old.claims_indexed.filter(start_byte__isnull=True).exists():
            return old,True
        # Legacy indexes have no offsets; rebuild once for per-claim retrieval.
    after=path.stat()
    if (before.st_size,before.st_mtime_ns)!=(after.st_size,after.st_mtime_ns):
        raise ValueError('EDI file changed during indexing; will retry next scan')
    # SHA256 computed incrementally, without loading huge remote EDI data into RAM.
    digest=hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b''):
            digest.update(chunk)
    end=path.stat()
    if (before.st_size,before.st_mtime_ns)!=(end.st_size,end.st_mtime_ns):
        raise ValueError('EDI file changed during checksum; will retry next scan')
    # A source file may move from out/ into out/backup hours later.
    # Content SHA-256, not filename/path, identifies an unchanged EDI file.
    # Keep one EdiFile + its claim indexes and repair the path if it moved.
    duplicate = EdiFile.objects.filter(
        client=client, kind=kind, fingerprint=digest.hexdigest()
    ).exclude(path=str(path)).order_by('id').first()
    if not force and duplicate:
        previous = Path(duplicate.path)
        if not previous.is_file():
            with transaction.atomic():
                # Do not overwrite another indexed record at the destination.
                occupied = EdiFile.objects.filter(
                    client=client, kind=kind, path=str(path)
                ).exclude(pk=duplicate.pk).exists()
                if not occupied:
                    duplicate.path = str(path)
                    duplicate.filename = path.name[:255]
                    duplicate.bytes_size = end.st_size
                    duplicate.modified_ns = end.st_mtime_ns
                    duplicate.save(update_fields=[
                        'path', 'filename', 'bytes_size', 'modified_ns'
                    ])
        # Existing EdiClaimIndex/EdiDocument relationships remain intact.
        # A scan of an identical copy must never add another claim index.
        return duplicate, True
    with transaction.atomic():
        obj,_=EdiFile.objects.update_or_create(client=client,kind=kind,path=str(path),defaults={
            'filename':path.name[:255],'fingerprint':digest.hexdigest(),'bytes_size':end.st_size,'modified_ns':end.st_mtime_ns})
        obj.claims_indexed.all().delete()
        pending=[]
        for x,t,n,offset in edi_claim_refs(path,kind):
            pending.append(EdiClaimIndex(edi_file=obj,claim_identifier=x,identifier_type=t,segment_number=n,start_byte=offset))
            if len(pending)>=1000:
                EdiClaimIndex.objects.bulk_create(pending,batch_size=1000)
                pending=[]
        if pending:
            EdiClaimIndex.objects.bulk_create(pending,batch_size=1000)
        verified=path.stat()
        if (before.st_size,before.st_mtime_ns)!=(verified.st_size,verified.st_mtime_ns):
            raise ValueError('EDI changed during DB indexing; rolled back transaction')
    # EDI indexing only stores permanent file and claim references.
    # Run reconciliation once after the indexing batch finishes.
    return obj,False

@transaction.atomic
def reconcile_client(client):
    # Scoped by client on BOTH sides. Ambiguous mappings should be reviewed instead of joined by guesswork.
    claims=list(Claim.objects.filter(client=client).only('id','claim_number','dcn','error_message'))
    by_id={};by_dcn={}
    for c in claims:
        by_id.setdefault((c.claim_number or '').strip().upper(),[]).append(c)
        if c.dcn:by_dcn.setdefault(c.dcn.strip().upper(),[]).append(c)
    history_pending=[]
    for o in ErrorObservation.objects.filter(client=client).order_by('report_date','id').iterator(chunk_size=1000):
        # The WGS inventory CLAIM_NUMBER corresponds to the report DCN.
        # Do not equate the report's separate CLAIM SUBMITTER ID with DCN.
        key=(o.dcn or '').strip().upper()
        matches=by_id.get(key,[]) if key else []
        if not matches:matches=by_dcn.get(key,[]) if key else []
        if len(matches)!=1:continue
        c=matches[0]
        history_pending.append(ErrorHistory(claim_id=c.id,observation_id=o.id,report_date=o.report_date,message=o.message))
        if len(history_pending)>=1000:
            ErrorHistory.objects.bulk_create(history_pending,batch_size=1000,ignore_conflicts=True)
            history_pending=[]
    if history_pending:
        ErrorHistory.objects.bulk_create(history_pending,batch_size=1000,ignore_conflicts=True)
    from django.db.models import OuterRef,Subquery,Value,TextField
    from django.db.models.functions import Coalesce
    latest=ErrorHistory.objects.filter(claim_id=OuterRef('pk')).order_by('-report_date','-id').values('message')[:1]
    Claim.objects.filter(client=client).update(error_message=Coalesce(Subquery(latest),Value(''),output_field=TextField()))
    # Rebuild only presentation links; permanent claim-level file indexes remain independent
    # of the current LAG. Each CLP/CLM occurrence is an individual selectable item.
    EdiDocument.objects.filter(claim__client=client).delete()
    linked=[]
    for index in EdiClaimIndex.objects.filter(edi_file__client=client).select_related('edi_file').iterator(chunk_size=1000):
        identifier=(index.claim_identifier or '').strip().upper()
        matches=[]

        if client.code=='WGS' and index.edi_file.kind=='835':
            if index.identifier_type!='payer_control':
                continue
            if len(identifier)!=15 or not identifier.startswith('20'):
                continue
            matches=by_id.get(identifier[2:13],[])
        elif (client.code == 'WGS' and index.edi_file.kind == '837'
              and index.identifier_type == 'wgs_ref_9c'):
            matches = by_id.get(identifier, [])
        elif index.identifier_type=='submitter':
            matches=by_id.get(identifier,[])

        if len(matches)!=1:
            continue

        linked.append(EdiDocument(claim=matches[0],edi_file=index.edi_file,edi_index=index,
                                  kind=index.edi_file.kind,original_name=index.edi_file.filename))
        if len(linked)>=1000:
            EdiDocument.objects.bulk_create(linked,batch_size=1000)
            linked=[]
    if linked:
        EdiDocument.objects.bulk_create(linked,batch_size=1000)
