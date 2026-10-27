from datetime import date
from django.contrib.auth import get_user_model
from django.test import TestCase, Client as HttpClient
from django.db import connection
from django.test.utils import CaptureQueriesContext
from rest_framework.authtoken.models import Token
from .models import Client, Claim, Membership, ImportBatch, LagClaimSnapshot, ErrorHistory

class FastClaimQueryTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.tenant=Client.objects.create(code='FAST',name='Fast')
        cls.other=Client.objects.create(code='OTHER',name='Other')
        user=get_user_model().objects.create_user(username='filter-test',password='test')
        Membership.objects.create(user=user,client=cls.tenant)
        cls.token=Token.objects.create(user=user)
        cls.batch=ImportBatch.objects.create(client=cls.tenant,source='lag',filename='test.xlsx',report_date=date(2026,8,3),status='processed',original_file=b'x'*1024*1024)
        cls.records=[]
        for i in range(12):
            claim=Claim.objects.create(client=cls.tenant,claim_number=f'C{i:03d}',report_date=date(2026,8,3))
            cls.records.append(claim)
            LagClaimSnapshot.objects.create(client=cls.tenant,batch=cls.batch,claim=claim,row_number=i,sheet_name='LOCAL CLMS',row_data={'_snapshot':{'claim_number':claim.claim_number,'fund':'Fund A' if i%2 else 'Fund B','tpa_age':i,'total_age':i+20}})
            if i%3:
                ErrorHistory.objects.create(claim=claim,report_date=date(2026,8,2),message='Error A' if i%2 else 'Error B')
            ErrorHistory.objects.create(claim=claim,report_date=date(2026,8,4),message='Future error')
        # A duplicate with different data must never replace the first row.
        LagClaimSnapshot.objects.create(client=cls.tenant,batch=cls.batch,claim=cls.records[0],row_number=100,sheet_name='LOCAL ADJ',row_data={'_snapshot':{'claim_number':'C000','fund':'Wrong duplicate','tpa_age':999,'total_age':999}})
        otherclaim=Claim.objects.create(client=cls.other,claim_number='PRIVATE',report_date=date(2026,8,3))
        otherbatch=ImportBatch.objects.create(client=cls.other,source='lag',filename='other.xlsx',report_date=date(2026,8,3),status='processed')
        LagClaimSnapshot.objects.create(client=cls.other,batch=otherbatch,claim=otherclaim,row_number=1,sheet_name='LOCAL CLMS',row_data={'_snapshot':{'claim_number':'PRIVATE','fund':'Private','tpa_age':9,'total_age':9}})

    def get(self,query='',endpoint='claims/'):
        return HttpClient().get('/api/'+endpoint+'?client=FAST&lag_date=2026-08-03'+query,HTTP_AUTHORIZATION='Token '+self.token.key)

    def test_open_and_closed_inclusive_age_bounds(self):
        response=self.get('&tpa_age_min=5')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['count'],7)
        self.assertEqual(self.get('&tpa_age_min=5&tpa_age_max=5').json()['count'],1)
        self.assertEqual(self.get('&total_age_max=23').json()['count'],4)
        self.assertEqual(self.get('&tpa_age_min=0&tpa_age_max=0').json()['count'],1)
        self.assertEqual(self.get('&tpa_age_min=8&tpa_age_max=2').status_code,400)
        self.assertEqual(self.get('&tpa_age_min=abc').status_code,400)

    def test_multiple_selection_or_within_and_between(self):
        self.assertEqual(self.get('&fund=Fund+A&fund=Fund+B').json()['count'],12)
        self.assertEqual(self.get('&error=Error+A&error=Error+B').json()['count'],8)
        self.assertEqual(self.get('&error=none&error=Error+A').json()['count'],8)
        data=self.get('&fund=Fund+A&error=Error+A&tpa_age_min=5').json()
        self.assertEqual(data['count'],3)

    def test_deduplication_history_cutoff_and_tenant_scope(self):
        data=self.get().json()
        self.assertEqual(data['count'],12)
        self.assertNotIn('Future error',[r['error_message'] for r in data['results']])
        self.assertNotIn('Wrong duplicate',[r['fund'] for r in data['results']])
        self.assertNotIn('PRIVATE',[r['claim_number'] for r in data['results']])
        self.assertEqual(self.get('&search=future').json()['count'],0)

    def test_page_sql_limits_and_no_workbook_blob(self):
        with CaptureQueriesContext(connection) as queries:
            response=self.get('&page_size=3&page=2')
        self.assertEqual(len(response.json()['results']),3)
        self.assertLessEqual(len(queries),8)
        sql='\n'.join(q['sql'] for q in queries)
        self.assertNotIn('"original_file"',sql)
        self.assertIn('LIMIT 3 OFFSET 3',sql)

    def test_dashboard_and_export_share_filters(self):
        response=self.get('&fund=Fund+A&fund=Fund+B',endpoint='dashboard/')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['total'],12)
        self.assertEqual(response.json()['with_errors'],8)
        self.assertEqual(len(response.json()['attention']),6)
        response=self.get('&tpa_age_min=5&tpa_age_max=5',endpoint='claims/export/')
        self.assertEqual(response.status_code,200)
        from openpyxl import load_workbook
        import io
        book=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(book.active.max_row,2)
